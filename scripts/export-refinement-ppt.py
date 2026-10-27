"""Export SVG primitives as editable PowerPoint shapes (no slide screenshots)."""
from pathlib import Path
import json, re, sys
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
out=Path(sys.argv[1] if len(sys.argv)>1 else '/workspace/artifacts/remap-refinement-20261027')
prs=Presentation(); prs.slide_width=Inches(18); prs.slide_height=Inches(10.8)
def color(v):
 v=v.lstrip('#'); v=''.join(x*2 for x in v) if len(v)==3 else v
 return RGBColor.from_string(v.upper())
def box(sl,x,y,w,h,fill='#fff',stroke=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100))
 if fill=='none':sh.fill.background()
 else:sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
 if stroke and stroke!='none': sh.line.color.rgb=color(stroke);sh.line.width=Pt(.72)
 else:sh.line.fill.background()
 return sh
def txt(sl,x,y,value,size=20,fill='#233c51',bold=False,width=1700,url=None):
 sh=sl.shapes.add_textbox(Inches(x/100),Inches((y-size*1.04)/100),Inches(width/100),Inches(size*1.8/100));tf=sh.text_frame
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
 p=tf.paragraphs[0];p.space_before=p.space_after=Pt(0);r=p.add_run();r.text=value;r.font.name='Noto Sans CJK SC';r.font.size=Pt(size*.72);r.font.bold=bold;r.font.color.rgb=color(fill)
 if url:r.hyperlink.address=url
 return sh
def line(sl,x,y,x2,y2,fill,width=1):
 sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x/100),Inches(y/100),Inches(x2/100),Inches(y2/100));sh.line.color.rgb=color(fill);sh.line.width=Pt(width*.72)
def new(title,subtitle):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);box(sl,0,0,1800,1080,'#f4f7f9');box(sl,0,0,1800,102,'#17384d');txt(sl,30,43,title,28,'#ffffff',True);txt(sl,30,78,subtitle,16,'#d5e3ec');return sl
def svg_slide(name):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);root=ET.parse(out/name).getroot()
 def walk(el,attrs=None,url=None):
  attrs={**(attrs or {}),**el.attrib};tag=el.tag.split('}')[-1]
  if tag=='rect':box(sl,float(attrs.get('x',0)),float(attrs.get('y',0)),float(attrs['width']),float(attrs['height']),attrs.get('fill','#fff'),attrs.get('stroke'))
  if tag=='text':txt(sl,float(attrs['x']),float(attrs['y']),''.join(el.itertext()),float(attrs.get('font-size',16)),attrs.get('fill','#233c51'),float(attrs.get('font-weight',400))>=600,url=url)
  if tag=='path':
   tokens=re.findall(r'[MLHVZmlhvz]|-?\d+(?:\.\d+)?',attrs['d']);i=0;x=y=0;command=''
   while i<len(tokens):
    if tokens[i].isalpha():command=tokens[i];i+=1
    if command in ('M','L'): xx,yy=map(float,tokens[i:i+2]);i+=2
    elif command=='H':xx=float(tokens[i]);yy=y;i+=1
    elif command=='V':xx=x;yy=float(tokens[i]);i+=1
    else:break
    if command!='M':line(sl,x,y,xx,yy,attrs.get('stroke','#397e85'),float(attrs.get('stroke-width',1)))
    x,y=xx,yy
   return
  for child in el:walk(child,attrs,el.attrib.get('href',url))
 walk(root);return sl
sl=new('ReMAP · 可靠性设计与评估','架构 / 损失 / 措施 / 分解分配 / 业界验证 · 2026-10-27')
for y,title,body in [(180,'评审重点','架构增加颜色层次；光模块与关键部件显式展示；三阶段六类损失映射到措施与参数。'),(320,'指标条件','600小时固定窗口、98.5%目标、9小时预算。MTBF按事件范围标注；等效MTTR含全部损失。'),(460,'快照边界','本环境缺失上一轮未提交文件与原快照；本包是当前版本重建的独立情景，不能视为原结果重放。'),(600,'计算与证据','现状/延长线128轨迹、突破1024轨迹；主种子7419231。完整输入、哈希与损失账见JSON。'),(740,'设计建议','先校准现场BOM和故障日志；再评估快速检测恢复、分级检查点及跨故障域冗余组合。'),(880,'辅助模型','网络BDD与20万次独立MC使用演示FIT；结构可用度和600小时训练可用度分别判断。')]:
 txt(sl,50,y,title,24,'#316b9a',True);txt(sl,50,y+46,body,21)
svg_slide('architecture.svg');svg_slide('loss.svg');svg_slide('scales.svg')
n=json.load(open(out/'network-study.json'))['full'];sl=new('组网辅助对照 · 光模块与共享故障','演示参数，非项目预测 · 相同容量条件 · BDD与独立蒙特卡洛')
txt(sl,40,165,'8 / 8平面必需；每路径2个光模块；互联柜共享事件、超节点内共享事件均显式建模。',23)
columns=[40,430,730,1050,1510]
headers=['方案','BDD结构可用度','MC估计','MC 95% Wilson区间','失败样本']
for x,h in zip(columns,headers):txt(sl,x,245,h,22,'#316b9a',True)
for i,r in enumerate(n['rows']):
 y=310+i*100;box(sl,30,y-35,1740,86,'#fff');sample=r['sample'];values=[r['name'],f"{r['availability']*100:.5f}%",f"{sample['availability']*100:.5f}%",f"{sample['interval95'][0]*100:.5f}–{sample['interval95'][1]*100:.5f}%",f"{sample['failures']} / {sample['trials']}"]
 for x,v in zip(columns,values):txt(sl,x,y,v,22)
for y,v in [(660,'L2：100,000 FIT / 8h；光模块：50,000 FIT / 4h；机柜：20,000 FIT / 4h。'),(715,'超节点内共享故障：10,000 FIT / 4h。原始器件及机柜原因必须互斥；不能重复扣除残余FABRIC保护收益。'),(795,'跨柜冗余能降低共享机柜风险，但仍受上游供电、控制、光路等公共依赖限制。'),(850,'7 / 8容量敏感性另见完整记录；容许少一个平面必须同时计入吞吐损失。'),(930,'结构辅助结果不自动改写主模型FIT；倒换瞬态、维修排队、流量重建需动态仿真。')]:txt(sl,40,y,v,21)
svg_slide('research.svg')
for cards in [10000,100000,200000]:svg_slide(f'loss-{cards}.svg')
prs.save(out/'ReMAP-架构与可靠性措施-九页可编辑.pptx')
assert len(prs.slides)==9
print('9 editable slides;',sum(len(s.shapes) for s in prs.slides),'native shapes')
