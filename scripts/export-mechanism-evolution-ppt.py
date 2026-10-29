"""Export editable ReMAP summary and detail slides from source/presentation-data.json.
Usage: python scripts/export-mechanism-evolution-ppt.py /path/to/review-directory
Dependencies: python-pptx, Pillow; Noto Sans CJK SC installed for measurement.
Charts and text use native PowerPoint shapes. Frozen source data is not edited.
"""
from pathlib import Path
import sys
from collections import defaultdict
import hashlib, json, math, re
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else Path("/workspace/artifacts/remap-mtbf-mttr-heatmaps-20261029")
D = json.loads((ROOT / 'source/presentation-data.json').read_text())
T = D['comparison']['topologyDecision']
S = next(s for s in T['scales'] if s['cards'] == 500000)
STAGES = S['scenarios']
NS = [{n['id']: n for n in s['nodes']} for s in STAGES]
IDS = ['domain-1','hardware-network','hardware-storage','hardware-facility','hardware-other','overhead']
DOMAINS = ['01 计算资源','02 网络互联','03 共享存储','04 设施与管理','05 软件与其他','06 公共开销']
LOSS = [('恢复停机',['downtime']),('回滚重算',['rollback']),('亚健康 / 局部降容',['slowdown','localLoss']),('污染失效',['terminalInvalid']),('保存阻塞',['save']),('检测开销',['mechanismOverhead'])]
C = dict(navy='18394F', blue='356D9A', teal='267F74', gold='B98322', ink='20384A', muted='5F7180', line='D5DFE6', pale='F3F6F8', white='FFFFFF', red='A4543B', lightblue='EAF1F8', lightteal='EAF5F1', lightgold='FBF2DC')
LOSS_COLORS = ['CF9B35','587CAA','6D9A91','AA7764','B9A261','8997A7']
FAMILY = 'Noto Sans CJK SC'
FONT_PATH = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FONT = ImageFont.truetype(FONT_PATH, 100)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(20), Inches(11.25)
prs.core_properties.title = 'ReMAP｜现状分解、损失与可靠性措施'
prs.core_properties.subject = '四部分方法与场景计算结果；独立验证与冻结目标'
prs.core_properties.author = 'ReMAP'
prs.core_properties.keywords = '可靠性,可用度,MTBF,MTTR,FIT,AGREE,Shapley,2026'
CHECKS = []
SLIDE_TEXT = []

def rgb(s):
    v=C.get(s,s).lstrip('#').upper()
    return RGBColor.from_string(''.join(c*2 for c in v) if len(v)==3 else v)
def width(text, size): return FONT.getlength(str(text)) / 100 * size / 72
def wrap(text, size, w):
    result=[]
    for raw in str(text).split('\n'):
        cur=''
        for token in re.findall(r'[A-Za-z0-9._/%+\-]+|.', raw):
            if cur and width(cur+token,size) > w-.025:
                result.append(cur.rstrip());cur=token.lstrip()
            else:cur+=token
        result.append(cur)
    return '\n'.join(result)

def rect(sl,x,y,w,h,fill='white',line=None,radius=False,lw=1):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line);sh.line.width=Pt(lw)
    else:sh.line.fill.background()
    if radius:
        try:sh.adjustments[0]=.08
        except (IndexError,AttributeError):pass
    return sh

def txt(sl,x,y,w,h,text,size=16,color='ink',bold=False,align=PP_ALIGN.LEFT,center=False,url=None,auto_wrap=False):
    value=wrap(text,size,w) if auto_wrap else str(text)
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    sh.name='Text: '+str(text).replace('\n',' / ')[:100]
    tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE if center else MSO_ANCHOR.TOP
    for i,s in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.space_before=p.space_after=Pt(0);p.line_spacing=1.45 if len(value.split("\n"))>1 else 1.1
        r=p.add_run();r.text=s;r.font.name=FAMILY;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        rp=r._r.get_or_add_rPr();ea=OxmlElement('a:ea');ea.set('typeface',FAMILY);rp.append(ea)
        if url:
            r.hyperlink.address=url;r.font.underline=False;r.font.color.rgb=rgb(color)
        if width(s,size)>w+.045:CHECKS.append({'slide':len(prs.slides),'kind':'width','text':s,'width':width(s,size),'box':w})
    estimated=size/72*(1.1+1.45*(len(value.split("\n"))-1))
    if estimated>h+.045:CHECKS.append({'slide':len(prs.slides),'kind':'height','text':value,'height':estimated,'box':h})
    if min(x,y)<-.005 or x+w>20.005 or y+h>11.255:CHECKS.append({'slide':len(prs.slides),'kind':'bounds','text':value})
    SLIDE_TEXT[-1].append(value)
    return sh

def line(sl,x1,y1,x2,y2,color='line',lw=1.5,arrow=False,dash=False):
    sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(lw)
    lp=sh.line._get_or_add_ln()
    if arrow:
        a=OxmlElement('a:tailEnd');a.set('type','triangle');lp.append(a)
    if dash:
        a=OxmlElement('a:prstDash');a.set('val','dash');lp.append(a)
    return sh

def path(sl, points, color='blue',lw=2,arrow=True,dash=False):
    for i,(a,b) in enumerate(zip(points,points[1:])):line(sl,*a,*b,color,lw,arrow and i==len(points)-2,dash)

def pill(sl,x,y,w,text,color='blue',fill='lightblue',size=12):
    rect(sl,x,y,w,.31,fill,radius=True);txt(sl,x+.06,y+.025,w-.12,.25,text,size,color,True,PP_ALIGN.CENTER)

def p(v,d=2):return '—' if v is None else f'{v*100:.{d}f}%'
def h(v,d=3):return '—' if v is None else '0' if abs(v)<1e-9 else '<0.001' if 0<v<.001 else f'{v:.{d}f}'
def tm(v):return '—' if v is None else f'{v:.2f} h' if v>=1 else f'{v*60:.2f} min'
def losses(n):return [sum(n['loss'][k] for k in keys) for _,keys in LOSS]
def heat(v):
    t=max(0,min(1,math.log1p(v*60)/math.log1p(600*60)))
    return ''.join(f'{round(a+(b-a)*t):02X}' for a,b in zip([255,241,243],[213,108,58]))
def table(sl,x,y,widths,headers,rows,rh=.56,hh=.6,size=14,colors=None,aligns=None):
    xx=x
    for w,head in zip(widths,headers):
        rect(sl,xx,y,w,hh,'navy');txt(sl,xx+.08,y+.045,w-.16,hh-.09,head,13,'white',True,center=True);xx+=w
    for i,row in enumerate(rows):
        xx=x;yy=y+hh+i*rh
        for j,(w,value) in enumerate(zip(widths,row)):
            bg='pale' if i%2==0 else 'white';fg='ink'
            if colors:
                opts=colors(i,j)
                if opts:bg,fg=opts
            rect(sl,xx,yy,w,rh,bg)
            txt(sl,xx+.08,yy+.045,w-.16,rh-.08,value,size,fg,j==0,align=aligns[j] if aligns else PP_ALIGN.LEFT,center=True)
            xx+=w
        line(sl,x,yy+rh,x+sum(widths),yy+rh,'line',.5)

import xml.etree.ElementTree as ET
V=json.loads((ROOT/'source/visual-data.json').read_text())
BM=V['matrix']; EV=next(s['evolution'] for s in V['scales'] if s['cards']==500000)
CUR,EXT,BR=[n['system'] for n in NS]
SHORT=['现状','延长线改进','突破性改进']; COLS=['blue','red','teal']
C.update(orange='D56C3A',pink='FFF1F3',gain='5999CD')
LOSS_COLORS[:]=['D17B53','6B91BA','79A49E','BB8D91','D9AB9F','9CABB8']
RANK=sorted(range(len(BM['rows'])),key=lambda i:-(BM['stages'][2]['cells'][i][3]['pp'] or 0))
RANKS={i:k+1 for k,i in enumerate(i for i in RANK if BM['stages'][2]['cells'][i][3]['pp'] is not None)}

def newslide(title,subtitle):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);SLIDE_TEXT.append([])
    rect(sl,0,0,20,11.25,'white');rect(sl,0,0,20,.065,'navy')
    txt(sl,.4,.21,19.15,.46,title,27,'navy',True)
    txt(sl,.42,.78,19.10,.26,subtitle,13.0,'muted')
    return sl

def footer(sl,text):txt(sl,.42,10.96,19.1,.2,text,10.0,'muted')
def section(sl,x,y,w,title,col='blue',size=14):
    rect(sl,x,y,.05,.25,col);txt(sl,x+.13,y-.015,w-.15,.29,title,size,col,True)
def flatten():
    for slide in prs.slides:
        for shape in slide.shapes:
            for effect in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):effect.set('idx','0')
            sp=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
            if sp is not None and sp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst') is None:sp.append(OxmlElement('a:effectLst'))

def gainfill(value):
    if value is None:return 'F1F3F5'
    t=math.sqrt(min(1,abs(value)/BM['ceiling']))
    a,b=([241,247,253],[89,153,205]) if value>=0 else ([255,241,243],[220,133,105])
    return ''.join(f'{round(u+(v-u)*t):02X}' for u,v in zip(a,b))
def smallhours(v):return '0' if abs(v)<1e-9 else '<.01' if 0<v<.01 else f'{v:.2f}'

def loss_matrix(sl,x,y,w,rh=.27,size=12.0,notes=False):
    label=1.55;gap=.18;gw=(w-label-2*gap)/3;total=1.08;cw=(gw-total)/6;hh=.34;sh=.27
    for k in range(3):
        gx=x+label+k*(gw+gap);rect(sl,gx,y,gw,hh,'FFF1EE','E4C7BE')
        txt(sl,gx+.04,y+.04,gw-.08,hh-.04,f"{SHORT[k]}   A {p(NS[k]['system']['availability'],3)}",size+.2,'red',True,PP_ALIGN.CENTER)
        for j,lab in enumerate(['停机','回滚','降速','污染','保存','检测']):txt(sl,gx+j*cw,y+hh+.035,cw,sh-.03,lab,size-.6,'muted',True,PP_ALIGN.CENTER)
        txt(sl,gx+6*cw,y+hh+.035,total,sh-.03,'合计/预算',size-1.3,'muted',True,PP_ALIGN.CENTER)
    for i,(id_,label_) in enumerate(list(zip(IDS,DOMAINS))+[('system','全系统')]):
        yy=y+hh+sh+i*rh;txt(sl,x+.04,yy+.035,label-.10,rh-.035,label_.replace('设施与管理','设施管理').replace('软件与其他','软件其他'),size-.4,'ink',i==6)
        for k in range(3):
            gx=x+label+k*(gw+gap);n=NS[k][id_]
            for j,v in enumerate(losses(n)):
                rect(sl,gx+j*cw,yy,cw-.018,rh-.018,heat(v));txt(sl,gx+j*cw,yy+.035,cw-.025,rh-.035,smallhours(v),size,'ink',i==6,PP_ALIGN.CENTER)
            txt(sl,gx+6*cw,yy+.035,total,rh-.035,f"{n['actualHours']:.2f}/{n['allowedHours']:.2f}",size-2.4,'red' if (n['gapHours'] or 0)>0 else 'teal',True,PP_ALIGN.CENTER)
    by=y+hh+sh+7*rh
    if notes:
        txt(sl,x,by+.05,w,.24,'M3/M4/M8→停机   M4/M5/M6→回滚   M7→降速   M6→污染   M5→保存；保存/检测开销归06。共用 log(1+小时/1分钟) 色尺。',11.6,'muted')
        by+=.32
    return by

def benefit_matrix(sl,x,y,w,rh=.26,size=12.0,total=False):
    label=3.40;gap=.20;gw=(w-label-2*gap)/3;cw=gw/4;hh=.53 if size>=13 else .44;sh=.25
    for k,stage in enumerate(BM['stages']):
        gx=x+label+k*(gw+gap);rect(sl,gx,y,gw,hh,'E9F2FA','BCD1E2')
        txt(sl,gx+.04,y+.025,gw-.08,.21,stage['name'],size+.3,'blue',True,PP_ALIGN.CENTER)
        txt(sl,gx+.04,y+(.31 if size>=13 else .24),gw-.08,.18,'已有贡献 / 原固定基线' if k==0 else '改进增量 / 冻结现状',size-2.0,'muted',False,PP_ALIGN.CENTER)
        for j,lab in enumerate(['万卡','十万卡','二十万卡','50万卡']):txt(sl,gx+j*cw,y+hh+.025,cw,sh-.02,lab,size-.7,'muted',False,PP_ALIGN.CENTER)
    labels=['MTBF · M1 故障源治理','MTBF · M2 冗余 / 副本恢复','MTTR · M3 中断检测','MTTR · M4/M8 恢复与升级','等效MTTR · M5 异步检查点','等效MTTR · M6 静默检出','等效MTTR · M6 漏检排查','等效MTTR · M7 亚健康检测']
    for i,(row,lab) in enumerate(zip(BM['rows'],labels)):
        yy=y+hh+sh+i*rh
        rank=str(RANKS[i]) if i in RANKS else '—'
        rect(sl,x,yy,.28,rh-.015,'lightblue');txt(sl,x,yy+.027,.28,rh-.025,rank,size-1,'blue',True,PP_ALIGN.CENTER)
        txt(sl,x+.36,yy+.025,label-.42,rh-.025,lab,size-.5,'ink')
        for k,stage in enumerate(BM['stages']):
            gx=x+label+k*(gw+gap)
            for j,c in enumerate(stage['cells'][i]):
                v=c['pp'];value='—' if v is None else '0' if abs(v)<1e-9 else f'{v:+.3f}' if abs(v)<.01 else f'{v:+.2f}'
                rect(sl,gx+j*cw,yy,cw-.025,rh-.015,gainfill(v));txt(sl,gx+j*cw,yy+.027,cw-.03,rh-.025,value,size,'muted' if v is None else 'ink',False,PP_ALIGN.CENTER)
    by=y+hh+sh+8*rh
    if total:
        txt(sl,x+.03,by+.035,label-.08,rh-.025,'本研究组合收益 / pp',size-.5,'blue',True)
        for k,stage in enumerate(BM['stages']):
            gx=x+label+k*(gw+gap)
            for j,v in enumerate(stage['gains']):
                rect(sl,gx+j*cw,by,cw-.025,rh-.015,'lightblue');txt(sl,gx+j*cw,by+.035,cw-.03,rh-.025,'—' if v is None else f'{v:+.3f}' if abs(v)<.01 else f'{v:+.2f}',size,'blue',True,PP_ALIGN.CENTER)
        by+=rh
    return by

def compact_topology(sl,x,y,w):
    # A simplified version of the platform reference: physical containment on left, network planes on right.
    rect(sl,x,y,w,.30,'EAF0FA','B8CBE1',True);txt(sl,x+.10,y+.04,w-.20,.23,'02 参数面：Leaf ↔ Spine   /   03 存储面',11.8,'blue',True)
    rect(sl,x,y+.41,w*.58,1.18,'F1F8F4','A1C6B6',True)
    txt(sl,x+.10,y+.46,w*.58-.2,.23,'01 超节点：16计算柜 × 64 NPU',11.4,'teal',True)
    rect(sl,x+.10,y+.76,w*.58-.20,.66,'white','BED6CA',True)
    txt(sl,x+.18,y+.79,w*.58-.36,.18,'代表性计算板 · 8 NPU',10.1,'teal')
    for i in range(8):
        xx=x+.18+i*(w*.58-.36)/8
        rect(sl,xx,y+1.05,(w*.58-.42)/8-.02,.19,'D4EADF','8FBDAA',True)
        txt(sl,xx,y+1.07,(w*.58-.42)/8-.02,.16,str(i+1),8.7,'teal',False,PP_ALIGN.CENTER)
    txt(sl,x+.13,y+1.43,w*.58-.24,.16,'CORE / SRAM / IOD / HBM · CPU / DDR / SSD',8.1,'muted')
    nx=x+w*.62;nw=w*.38
    rect(sl,nx,y+.41,nw,.49,'EAF0FA','A7BEDB',True);txt(sl,nx+.08,y+.46,nw-.16,.2,'L1 ↔ L2 · 8平面',11.4,'blue',True);txt(sl,nx+.08,y+.71,nw-.16,.17,'2互联柜 / 16 L2',9.7,'muted')
    rect(sl,nx,y+1.02,nw,.57,'FFF1ED','D5A18B',True);txt(sl,nx+.08,y+1.08,nw-.16,.20,'光模块 · 2+2保护',11,'red',True);txt(sl,nx+.08,y+1.34,nw-.16,.17,'残余故障已计入模型',9.7,'muted')
    path(sl,[(x+w*.58,y+.99),(nx-.07,y+.99),(nx-.07,y+.66),(nx,y+.66)],'blue',1.8,False)
    path(sl,[(nx+nw*.5,y+.9),(nx+nw*.5,y+1.02)],'blue',1.8,False)
    txt(sl,x,y+1.68,w,.21,'04 供电冷却   ·   05 软件控制   ·   06 保存/检测开销',10.8,'muted')

def svg_native(sl,file,x,y,w,max_y=None):
    root=ET.parse(file).getroot();scale=w/1920
    def walk(e):
        tag=e.tag.split('}')[-1];a=e.attrib
        if tag=='rect':
            xx=float(a.get('x',0));yy=float(a.get('y',0));ww=float(a.get('width',0));hh=float(a.get('height',0))
            if max_y is None or yy+hh<=max_y:rect(sl,x+xx*scale,y+yy*scale,ww*scale,hh*scale,a.get('fill','#ffffff'),a.get('stroke') if a.get('stroke') not in [None,'none'] else None,bool(float(a.get('rx',0))),.6)
        elif tag=='text':
            xx=float(a.get('x',0));yy=float(a.get('y',0));value=''.join(e.itertext());fs=float(a.get('font-size',13))*scale*72
            if max_y is not None and yy>max_y:return
            width_=min(20-x-xx*scale-.08,width(value,fs)+.13)
            txt(sl,x+xx*scale,y+yy*scale-fs/72*1.02,width_,fs/72*1.28,value,fs,a.get('fill','#20384a'),int(a.get('font-weight','400'))>=600)
        elif tag=='path':
            if a.get('stroke') in [None,'none']:return
            tokens=re.findall(r'[A-Za-z]|-?\d+(?:\.\d+)?',a['d']);i=0;pos=(0,0);cmd='M'
            while i<len(tokens):
                if tokens[i].isalpha():cmd=tokens[i];i+=1
                if cmd in ['M','L']:
                    dest=(float(tokens[i]),float(tokens[i+1]));i+=2
                    if cmd=='L':line(sl,x+pos[0]*scale,y+pos[1]*scale,x+dest[0]*scale,y+dest[1]*scale,a['stroke'],float(a.get('stroke-width','1.5'))*scale*72,dash='stroke-dasharray' in a)
                    pos=dest;cmd='L'
                elif cmd in ['H','V']:
                    dest=(float(tokens[i]),pos[1]) if cmd=='H' else (pos[0],float(tokens[i]));i+=1
                    line(sl,x+pos[0]*scale,y+pos[1]*scale,x+dest[0]*scale,y+dest[1]*scale,a['stroke'],float(a.get('stroke-width','1.5'))*scale*72,dash='stroke-dasharray' in a);pos=dest
                else:raise ValueError('Unsupported SVG path '+cmd)
        for child in e:walk(child)
    walk(root)

def recovery_bars(sl,x,y,w,rh=.7):
    label=1.6;val=1.7;bw=w-label-val;maxv=max(s['total'] for s in EV);cols=['6B94B8','367CAB','C68A70']
    for i,s in enumerate(EV):
        yy=y+i*rh;txt(sl,x,yy,label-.1,.29,SHORT[i],14.3,COLS[i],True)
        rect(sl,x+label,yy+.03,bw,.25,'pale');xx=x+label
        for v,col in zip([s['detection'],s['attempt'],s['escalation']],cols):
            width_=bw*v/maxv
            if width_>0:rect(sl,xx,yy+.03,width_,.25,col);xx+=width_
        txt(sl,x+label+bw+.10,yy,val-.10,.29,f"{s['total']:.3f} s",14.3,'blue',True,PP_ALIGN.RIGHT)
        txt(sl,x+label,yy+.35,bw+val,.24,f"检测 {s['detection']:.3f} + 恢复 {s['attempt']:.1f} + 失败升级期望 {s['escalation']:.1f} s",11.8,'muted')

sl=newslide('ReMAP｜先定位损失，再按收益选择措施，最后落到参数','50万卡 / 600 h · 目标98.5% / 9 h预算 · 粉橙：损失小时 · 蓝色：措施收益pp · 主结果沿用独立计算快照')
for i,(name,method,col) in enumerate([('01 分解分配','AGREE · 六域冻结预算','blue'),('02 评估差距','DES损失账 / RBD、BDD结构辅助','blue'),('03 机制设计','Shapley · 四规模收益归因','blue'),('04 机制验证','独立1024轨迹 · 工程参数验收','teal')]):
    x=.42+i*4.84;rect(sl,x,1.12,4.67,.48,'pale');txt(sl,x+.1,1.17,1.55,.24,name,12.3,col,True);txt(sl,x+1.65,1.19,2.9,.22,method,9.5,'muted')
section(sl,.42,1.78,6.3,'950参考架构 · 包含与连接',size=13.6)
compact_topology(sl,.48,2.12,6.1)
section(sl,6.85,1.78,8.08,'六层分配：现状A → 目标；突破仍超预算的位置',size=13.6)
rows=[]
for i,id_ in enumerate(IDS):
    n=NS[0][id_];b=NS[2][id_];g=b['gapHours'];rows.append([DOMAINS[i].replace('设施与管理','设施管理'),p(n['availability'],2)+' → '+p(n['target'],3),f"{n['allowedHours']:.3f}",f"+{g:.3f}" if g>0 else '预算内'])
table(sl,6.85,2.13,[1.73,3.05,1.18,1.42],['责任域','现状 A → 目标 A','预算/h','突破缺口/h'],rows,rh=.255,hh=.28,size=10.8)
rect(sl,15.45,1.80,4.12,2.16,'EAF5F1','B5D5C9',True)
txt(sl,15.62,1.94,3.77,.25,'突破独立验证 · 均值达标',14.3,'teal',True)
txt(sl,15.62,2.31,3.77,.62,p(BR['availability'],3),34,'teal',True)
st=STAGES[-1]['statistics']
txt(sl,15.62,3.02,3.77,.25,f"95%区间 {p(st['ciLow'],3)}–{p(st['ciHigh'],3)}",11.3,'teal')
txt(sl,15.62,3.35,3.77,.25,f"损失 {BR['actualHours']:.3f} / 9 h",14.0,'ink',True)
txt(sl,15.62,3.68,3.77,.2,'均值区间 ≠ 单窗口达标概率',10.3,'muted')
section(sl,.42,4.13,19.1,'02 损失热力图 / h：同一责任层横向对照三阶段；停机、回滚、降速等分列',col='orange',size=13.3)
loss_matrix(sl,.42,4.49,19.14,.235,11.4)
section(sl,.42,6.90,19.1,'03 措施收益热力图 / pp：每块按四种规模展开；左侧数字为50万卡突破收益排序',size=13.3)
benefit_matrix(sl,.42,7.24,19.14,.252,11.7)
txt(sl,.44,10.00,19.1,.23,'现状与改进的归因基线不同，不能跨块相加或相减；0=该基线下无增量，—=未单独评估。三块共用蓝色色尺；正值为减损。',10.6,'muted')
rect(sl,.42,10.32,19.14,.47,'lightblue')
txt(sl,.56,10.39,18.85,.30,'04 参数验收：单次中断期望链 352 → 132.4 → 19.718 s；全损失等效MTTR 1571.23 → 220.53 → 24.16 s；全系统MTBF 15.12 → 15.12 → 30.49 min',12.0,'blue',True)
footer(sl,'建议优先验证 M6静默检出、M5异步检查点、M4/M8恢复链，并补足计算域缺口。M2冗余/备用/副本恢复另做BDD与DES；0秒保存为假设。2026-10-29')
sl.notes_slide.notes_text_frame.text=json.dumps({'provenance':D['provenance'],'sourceSha256':V['sourceSha256'],'benefitMatrix':BM,'recoveryEvolution':EV,'statistics':st},ensure_ascii=False)
flatten();prs.save(ROOT/'ReMAP-六层分配与双热力图-一页.pptx')

sl=newslide('ReMAP｜沿用950参考架构，将六层目标与缺口写到图旁','01 AGREE：由98.5%拆为9 h预算，按复杂度、重要度与设计权重逐级分配；各规模独立分配、三阶段沿用冻结预算')
svg_native(sl,ROOT/'source/topology-500000-current.svg',.4,1.15,19.2,954)
txt(sl,.50,10.72,18.9,.23,'突破仍有缺口：01计算资源超预算3.755 h；05软件与其他超0.082 h。分层A为损失责任投影，不可相乘；详细分配、其他规模与部件证据见平台。',11.7,'red',True)
footer(sl,'代表性实例：1024 NPU / 超节点、16计算柜、8交换平面；采用2互联柜×8台L2，原图4柜口径差异需按BOM核对。2+2光保护已包含在残余故障模型。')

sl=newslide('ReMAP｜损失定位 → 措施收益：颜色、编号与基线一一对应','02 评估：六层 × 六类损失 / h · 03 设计：八项归因因子 × 四种规模 / pp · 两图均按现状、延长线、突破分三块')
section(sl,.42,1.24,19.1,'损失账：粉色低损失 → 橙色高损失；三阶段采用相同对数色尺',col='orange',size=15.2)
loss_matrix(sl,.42,1.65,19.14,.32,13.4,True)
section(sl,.42,4.95,19.1,'措施收益：蓝色正收益；三阶段四规模共同色尺；左侧序号为50万卡突破收益排名',size=15.2)
benefit_matrix(sl,.42,5.36,19.14,.36,14.0,True)
txt(sl,.44,9.45,19.1,.28,'基线：现状贡献相对本轮固定基线；延长线和突破增量相对冻结现状。后两组均为64条配对轨迹 × 32个排列，不能与主表差值强行相等。',12.3,'muted')
txt(sl,.44,9.85,19.1,.28,'推荐顺序（50万卡）：M6 静默检出 +18.88 pp → M5 异步检查点 +13.66 pp → M4/M8 恢复与升级 +10.44 pp → M1 FIT治理 +9.66 pp。',13.5,'blue',True)
txt(sl,.44,10.30,19.1,.28,'收益排序辅助选择措施；区间重叠时不强判。M2冗余/副本级恢复尚未独立计量；应检查光链路、共因、备用资源和训练副本边界。',12.3,'muted')
footer(sl,'0为该基线下没有边际增量；—为没有独立证据。蓝色共用√(pp/25)色尺，格内为真实pp；粉橙损失共用log(1+小时/1分钟)色尺。完整区间见平台悬停提示。')

sl=newslide('ReMAP｜把机制落实为时间变化与可验证的工程参数','04 机制验证：先用输入参数分解中断链，再用独立DES核查全部损失；工程恢复耗时与全损失等效MTTR分别展示')
section(sl,.42,1.28,11.3,'降低MTTR：检测 + 恢复尝试 + 失败升级期望',size=16)
recovery_bars(sl,.48,1.85,11.15,.85)
rect(sl,12.15,1.28,7.4,3.1,'EAF1F8','BCD0DF',True)
txt(sl,12.34,1.45,7.02,.31,'352 → 19.718 s，单次期望链减少94.398%',18.2,'blue',True)
txt(sl,12.34,1.97,7.02,1.18,'突破：检测 1.118 s + 恢复 15 s + 升级期望 3.6 s\n检测 = 99.8%×1 s + 0.2%×60 s\n失败升级 = 0.1%×3600 s',14.0,'ink')
txt(sl,12.34,3.28,7.02,.72,'以上不含回滚、停机重叠与窗口裁剪；\n人工兜底实际追加1小时，不能理解为3.6秒。',12.8,'muted')
headers=['机制 / 生效参数','现状','延长线改进','突破性改进']
rows=[['M1 全系统MTBF / min']+[f"{s['mtbfMinutes']:.3f}" for s in EV],['M3 中断检出率 / 检测秒']+[f"{p(s['coverage'],2)} / {s['detect']:g} s" for s in EV],['M4 恢复尝试 / 成功率']+[f"{s['attempt']:g} s / {p(s['success'],2)}" for s in EV],['M8 失败时额外人工时间']+[f"{s['manual']/3600:g} h" for s in EV],['M5 检查点周期 / 保存阻塞']+[f"{s['checkpoint']:g} s / {s['save']:g} s" for s in EV],['M6 静默检出率 / 检测 / 漏检排查']+[f"{p(s['sdcCoverage'],2)} / {s['sdcDetect']:g} s / {s['sdcDelayHours']:g} h" for s in EV],['M7 亚健康检出率 / 检测秒']+[f"{p(s['slowCoverage'],2)} / {s['slowDetect']:g} s" for s in EV]]
table(sl,.45,4.65,[5.2,4.6,4.6,4.6],headers,rows,rh=.40,hh=.42,size=14)
for i,s in enumerate(EV):
    x=.45+i*6.45;rect(sl,x,8.13,6.28,1.02,'pale','line',True);txt(sl,x+.15,8.23,5.98,.24,SHORT[i]+' · 全损失等效MTTR',13,'muted');txt(sl,x+.15,8.59,2.86,.38,f"{s['equivalentMttrSeconds']:.3f} s",23,COLS[i],True);txt(sl,x+3.1,8.67,2.98,.25,f"条件预算 {s['targetMttrSeconds']:.3f} s",12,'muted')
txt(sl,.48,9.40,19.05,.28,'全损失等效MTTR = MTBF × (1−A) / A；包含停机、回滚、降速、污染、保存与检测开销，由损失账换算，不能替代物理维修时间。',13,'ink')
txt(sl,.48,9.87,19.05,.28,'新增机制建议：光链路多路径 / 备用池 / 副本级恢复，先用RBD/BDD核查结构与共因，再用相同故障样本DES评估备用耗尽及有效GPU·小时。',12.5,'blue',True)
txt(sl,.48,10.32,19.05,.27,'验收重点：99.91%静默检出、15秒/99.9%恢复、1秒检查点；实测持久化滞后、恢复带宽、误隔离、失败尾部与状态正确性。',12.6,'red',True)
footer(sl,'验证结果：50万卡均值98.696%，95%均值区间98.660%–98.733%，887/1024窗口达标。参数为情景要求，保存阻塞0秒为假设；现状基线仍需实测校准。')
flatten();prs.save(ROOT/'ReMAP-六层分配与双热力图-总分四页.pptx')
(ROOT/'source/ppt-layout-checks.json').write_text(json.dumps(CHECKS,ensure_ascii=False,indent=2))
(ROOT/'source/ppt-slide-text.json').write_text(json.dumps(SLIDE_TEXT,ensure_ascii=False,indent=2))
print(json.dumps({'slides':len(prs.slides),'native_shapes':[len(sl.shapes) for sl in prs.slides],'layout_warnings':CHECKS},ensure_ascii=False))
